"""Editable PPTX (2 slides per keyword, 16:9) from decks in data/trends.json, built with native shapes/text via python-pptx.

python -m render.build_pptx [--out site/downloads] [--only ENTITY_ID]
Fonts: Yu Gothic (East Asian) is requested; PowerPoint substitutes if absent. Text lengths are bounded by schema/trend.schema.json (deck).
"""
import argparse
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from collect.common import ROOT

PX = 9525  # EMU per CSS pixel at 96 dpi (1280px -> 13.333in)
FONT = "Yu Gothic"
NAVY, INK, MUTED, PAPER, WHITE, AMBER = "14213D", "1B2333", "5B6578", "F7F4EC", "FFFFFF", "FCA311"
ACCENT = {"model": "7A3DB8", "dev-tool": "0B7285", "local-runtime": "C2410C", "hardware": "2F6F3E", "research": "A1235B", "github-trending": "475569"}
VER = {"verified": ("検証済み", "157347"), "partially_verified": ("一部検証", "B26A00"), "unverified": ("未検証", "B02A37")}


def rgb(h):
    return RGBColor.from_string(h)


def px(v):
    return Emu(int(v * PX))


def rect(sl, x, y, w, h, fill=None, line=None, lw=0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = sl.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw * 0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _font(run, size, color, bold=False):
    run.font.size = Pt(size * 0.75)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(sl, x, y, w, h, paras, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0, margin=0):
    """paras: str | list of str | list of list[(text, {bold,color})] runs. `code` spans (backticks) use Consolas."""
    tb = sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = px(margin)
    tf.margin_top = tf.margin_bottom = px(margin)
    items = [paras] if isinstance(paras, str) else paras
    for i, p in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space * 0.75)
        runs = [(p, {})] if isinstance(p, str) else p
        for t, o in runs:
            for k, part in enumerate(re.split(r"`([^`]+)`", t)):
                if not part:
                    continue
                r = para.add_run()
                r.text = part
                _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold))
                if k % 2:
                    r.font.name = "Consolas"
    return tb


def head(sl, t, n, title, accent):
    rect(sl, 0, 0, 1280, 104, NAVY)
    x = 44
    chips = [(t["category"], accent, WHITE)] + ([(t["vendor"], "2C3E66", WHITE)] if t.get("vendor") else []) + \
            [("今週の注目" if t.get("trend_label") == "baseline" else "上昇中", AMBER, NAVY)]
    for label, bg, fg in chips:
        w = 26 + sum(15 if ord(ch) > 255 else 9 for ch in label)
        r = rect(sl, x, 12, w, 30, bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=15)
        text(sl, x, 12, w, 30, label, 16, fg, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        x += w + 10
    text(sl, 700, 12, 536, 30, f"{t['keyword']} ・ {n}/2", 16, "C9D1E3", False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    text(sl, 44, 48, 1192, 50, title, 34, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)


def foot(sl, t, refs):
    rect(sl, 44, 652, 1192, 2, "E3DCC8")
    src = [s for s in t["deck"]["sources"] if s["id"] in refs]
    lines = [[(f"[{s['id']}] ", {"bold": True}), (f"{s['title']}  ", {}), (s["url"], {"color": "1D4ED8"}), (f"  ({s['date']})" if s.get("date") else "", {})] for s in src]
    text(sl, 44, 658, 1000, 56, lines, 13, MUTED)
    label, col = VER[(t.get("verification") or {}).get("status", "unverified")]
    text(sl, 1040, 658, 196, 52, [[(label, {"bold": True, "color": col, "size": 16})], [("要約は自動生成+一次情報で確認", {"size": 11})]], 11, MUTED, align=PP_ALIGN.RIGHT)


def slide1(prs, t):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = rgb(PAPER)
    a, o, s0 = ACCENT.get(t["category"], "475569"), t["deck"]["overview"], (t.get("slides") or [{}])[0]
    head(sl, t, 1, s0.get("title") or t["keyword"], a)
    rect(sl, 44, 120, 1192, 62, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=8)
    rect(sl, 44, 120, 8, 62, a)
    text(sl, 66, 120, 1160, 62, o["takeaway"], 23, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    # column 1: stat + flow
    rect(sl, 44, 200, 330, 150, WHITE, a, 3, MSO_SHAPE.ROUNDED_RECTANGLE, 18)
    n = len(o["stat"]["value"])
    text(sl, 54, 214, 310, 70, o["stat"]["value"], 44 if n <= 8 else 34 if n <= 12 else 26, a, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(sl, 54, 288, 310, 56, o["stat"]["label"], 17, MUTED, False, PP_ALIGN.CENTER)
    text(sl, 44, 362, 330, 24, "しくみ", 16, MUTED, True, PP_ALIGN.CENTER)
    y = 390
    for i, st in enumerate(o["steps"]):
        rect(sl, 44, y, 330, 44, WHITE, a, 2, MSO_SHAPE.ROUNDED_RECTANGLE, 12)
        text(sl, 52, y, 314, 44, st, 16, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        y += 44
        if i < len(o["steps"]) - 1:
            text(sl, 44, y, 330, 14, "▼", 12, a, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            y += 14
    # column 2: changes
    text(sl, 392, 198, 516, 28, "何が変わる・何が起きた", 20, a, True)
    ch = o["changes"]
    h = (652 - 232 - 12 * (len(ch) - 1) - 8) / len(ch)
    for i, c in enumerate(ch):
        yy = 232 + i * (h + 12)
        rect(sl, 392, yy, 516, h, WHITE, a, 3, MSO_SHAPE.ROUNDED_RECTANGLE, 16)
        circ = rect(sl, 406, yy + h / 2 - 19, 38, 38, a, shape=MSO_SHAPE.OVAL)
        text(sl, 406, yy + h / 2 - 19, 38, 38, str(i + 1), 20, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(sl, 456, yy + 6, 440, h - 12, c, 20, INK, anchor=MSO_ANCHOR.MIDDLE)
    # column 3: facts
    text(sl, 926, 198, 310, 28, "要点", 20, a, True)
    yy = 232
    for f in o["facts"]:
        rect(sl, 926, yy, 104, 84, a, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=10)
        text(sl, 930, yy, 96, 84, f["k"], 17, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        rect(sl, 1030, yy, 206, 84, WHITE, a, 2, MSO_SHAPE.ROUNDED_RECTANGLE, 10)
        text(sl, 1038, yy, 192, 84, f["v"], 17, INK, False, anchor=MSO_ANCHOR.MIDDLE)
        yy += 96
    foot(sl, t, o["refs"])


def box(sl, x, y, w, h, title, items, a, kind="ul", warn=False):
    rect(sl, x, y, w, h, "FFFAF0" if warn else WHITE, "C98A00" if warn else a, 3, MSO_SHAPE.ROUNDED_RECTANGLE, 16)
    text(sl, x + 16, y + 8, w - 32, 28, title, 20, "8A5A00" if warn else a, True)
    paras = []
    for i, it in enumerate(items):
        if kind == "dl":
            paras.append([(it["t"] + "  ", {"bold": True, "color": a}), (it["d"], {})])
        else:
            paras.append([((f"{i + 1}. " if kind == "ol" else "• "), {"bold": True, "color": a}), (it, {})])
    text(sl, x + 16, y + 40, w - 32, h - 46, paras, (20 if sum(len(x) if isinstance(x, str) else len(x["d"]) for x in items) <= 180 else 18) if kind != "dl" else 18, INK, space=4)


def slide2(prs, t):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = rgb(PAPER)
    a, m = ACCENT.get(t["category"], "475569"), t["deck"]["impact"]
    head(sl, t, 2, f"{t['keyword']}: 誰に何が嬉しいか・使い方・注意点", a)
    rect(sl, 44, 120, 1192, 56, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=8)
    rect(sl, 44, 120, 8, 56, a)
    text(sl, 66, 120, 1160, 56, [[("想定する人   ", {"bold": True, "color": a}), (m["audience"], {})]], 22, INK, anchor=MSO_ANCHOR.MIDDLE)
    bw, bh = 587, 218
    box(sl, 44, 192, bw, bh, "嬉しいこと", m["benefits"], a)
    box(sl, 44 + bw + 18, 192, bw, bh, "使い方・始め方", m["how_to"], a, "ol")
    box(sl, 44, 192 + bh + 14, bw, bh, "注意点・制限", m["caveats"], a, warn=True)
    box(sl, 44 + bw + 18, 192 + bh + 14, bw, bh, "用語メモ", m["terms"], a, "dl")
    foot(sl, t, m["refs"])


def build(t, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1280), px(720)
    slide1(prs, t)
    slide2(prs, t)
    prs.core_properties.title = f"{t['keyword']} - AI Trend Radar"
    prs.core_properties.author = "AI Trend Radar"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))


def run(out="site/downloads", only=None):
    data = json.loads((ROOT / "data" / "trends.json").read_text(encoding="utf-8"))
    made = []
    for t in data["trends"]:
        if t.get("deck") and (not only or t["entity_id"] == only):
            build(t, ROOT / out / f"{t['entity_id']}.pptx")
            made.append(t["entity_id"])
    print(len(made), "PPTX ->", ROOT / out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="site/downloads")
    ap.add_argument("--only")
    a = ap.parse_args()
    run(a.out, a.only)


if __name__ == "__main__":
    main()
