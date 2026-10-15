"""Deck artifacts: PPTX structure and the HTML renderer's overflow-free layout for every committed deck."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from render import build_pptx  # noqa: E402

DATA = json.loads((ROOT / "data" / "trends.json").read_text(encoding="utf-8"))
WITH_DECK = [t for t in DATA["trends"] if t.get("deck")]


class PptxTests(unittest.TestCase):
    def test_every_deck_builds_two_editable_slides_with_sources(self):
        self.assertGreaterEqual(len(WITH_DECK), 5)
        with tempfile.TemporaryDirectory() as d:
            for t in WITH_DECK:
                out = Path(d) / f"{t['entity_id']}.pptx"
                build_pptx.build(t, out)
                prs = Presentation(str(out))
                self.assertEqual(2, len(prs.slides), t["entity_id"])
                text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
                self.assertIn(t["deck"]["overview"]["takeaway"], text)
                for s in t["deck"]["sources"]:
                    if s["id"] in t["deck"]["overview"]["refs"]:
                        self.assertIn(s["url"], text)

    def test_every_published_trend_has_digest_deck_and_one_infographic(self):
        for t in WITH_DECK:
            self.assertTrue(t.get("digest"), t["entity_id"])
            self.assertEqual(1, len(t["slides"]), t["entity_id"])
            self.assertEqual("verified", t["verification"]["status"].replace("partially_verified", "verified"))


class RenderTests(unittest.TestCase):
    def test_no_overflow_in_html_renderers(self):
        try:
            from playwright.sync_api import sync_playwright
        except ImportError:
            self.skipTest("playwright not installed")
        from render.build_site import OVERFLOW_JS, SITE
        bad = []
        with sync_playwright() as p:
            b = p.chromium.launch()
            page = b.new_page(viewport={"width": 1280, "height": 720})
            for t in WITH_DECK:
                for n in (0, 1):
                    page.goto((SITE / "deck.html").as_uri() + f"?e={t['entity_id']}&n={n}")
                    page.wait_for_selector("body[data-ready='1']")
                    bad += [f"{t['entity_id']} p{n + 1} {c}" for c in page.evaluate(OVERFLOW_JS)]
            b.close()
        self.assertEqual([], bad)


if __name__ == "__main__":
    unittest.main()
